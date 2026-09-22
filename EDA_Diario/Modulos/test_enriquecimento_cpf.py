"""
Testes do enriquecimento por CSV de CPF (sem MySQL).
Executar: python3 EDA_Diario/Modulos/test_enriquecimento_cpf.py
"""
from __future__ import annotations

import sys
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from modulo_enriquecimento_cpf import (  # noqa: E402
    EnriquecimentoCpfErro,
    classificar_blacklist_caso,
    gerar_excel_enriquecimento,
    ler_cpfs_csv,
    status_linha_sms,
    _linha_processo,
)


def _assert(cond: bool, msg: str) -> None:
    if not cond:
        raise AssertionError(msg)


def _bl_vazio() -> dict[str, set]:
    return {
        "CPF": set(),
        "NOME": set(),
        "TELEFONE": set(),
        "EMAIL": set(),
        "PROCESSO_INCIDENTE": set(),
    }


def test_csv_coluna_cpf() -> None:
    data = "CPF\n123.456.789-09\n12345678909\n987.654.321-00\n"
    cpfs = ler_cpfs_csv(data.encode("utf-8"))
    _assert(cpfs == ["12345678909", "98765432100"], f"unicidade/normalização: {cpfs}")


def test_csv_so_uma_coluna() -> None:
    data = "11144477735\n22255588846\n"
    cpfs = ler_cpfs_csv(data)
    _assert(cpfs == ["11144477735", "22255588846"], f"coluna única: {cpfs}")


def test_csv_ponto_virgula() -> None:
    data = "nome;CPF\nFulano;123.456.789-09\nBeltrano;11144477735\n"
    cpfs = ler_cpfs_csv(data)
    _assert(cpfs == ["12345678909", "11144477735"], f"CSV ; : {cpfs}")


def test_csv_vazio() -> None:
    try:
        ler_cpfs_csv("CPF\n")
    except EnriquecimentoCpfErro:
        return
    raise AssertionError("CSV sem CPF deveria falhar")


def test_blacklist_sim_motivo() -> None:
    bl = _bl_vazio()
    bl["CPF"].add("12345678909")
    out = classificar_blacklist_caso(
        cpf="123.456.789-09",
        nome="Fulano",
        requerente="Fulano",
        processo="0001234-56.2023.8.26.0100",
        incidente="1",
        telefones=["11999990000"],
        emails=["a@b.com"],
        bl=bl,
        motivo_map={("CPF", "12345678909"): "SEM INTERESSE"},
        data_map={("CPF", "12345678909"): datetime(2026, 8, 10, 14, 30)},
    )
    _assert(out["Blacklist"] == "Sim", f"devia ser Sim: {out}")
    _assert("SEM INTERESSE" in out["Motivo_blacklist"], f"motivo: {out}")
    _assert("CPF" in out["Tipos_blacklist"], f"tipos: {out}")
    _assert(out["Data_inclusao_blacklist"] == "10/08/2026 14:30", out["Data_inclusao_blacklist"])


def test_blacklist_nao() -> None:
    out = classificar_blacklist_caso(
        cpf="12345678909",
        nome="Fulano",
        requerente="Fulano",
        processo="PROC-1",
        incidente="1",
        telefones=["11988887777"],
        emails=[],
        bl=_bl_vazio(),
        motivo_map={},
    )
    _assert(out["Blacklist"] == "Não", f"devia ser Não: {out}")
    _assert(out["Motivo_blacklist"] == "", f"motivo vazio: {out}")


def test_excel_tem_coluna_blacklist() -> None:
    linhas = [
        {
            "CPF": "12345678909",
            "Requerente": "Fulano da Silva",
            "Numero_de_Processo": "PROC-1",
            "Numero_do_Incidente": "1",
            "Natureza": "",
            "Assunto": "",
            "Ordem": "",
            "Foro": "",
            "Principal_Liquido": None,
            "Calculo_Atualizado": None,
            "Entidade_Devedora": "",
            "Advogado": "",
            "Blacklist": "Sim",
            "Motivo_blacklist": "SEM INTERESSE",
            "_tels": ["11999990000"],
            "_emails": [],
            "_hsm": [],
        }
    ]
    buf = gerar_excel_enriquecimento(
        linhas,
        [{"CPF": "00000000000", "Blacklist": "Não", "Motivo_blacklist": ""}],
        {"CPFs_enviados": 2, "Processos_enriquecidos": 1},
    )
    import openpyxl

    wb = openpyxl.load_workbook(buf)
    _assert("Principal" in wb.sheetnames, wb.sheetnames)
    _assert("sms" in wb.sheetnames, wb.sheetnames)
    ws = wb["Principal"]
    headers = [c.value for c in ws[1]]
    _assert("TELEFONE_1" in headers, headers)
    _assert("Blacklist" in headers, headers)
    _assert("Motivo_blacklist" in headers, headers)
    _assert("Data_inclusao_blacklist" in headers, headers)
    bl_idx = headers.index("Blacklist") + 1
    _assert(ws.cell(2, bl_idx).value == "Sim", ws.cell(2, bl_idx).value)
    _assert(ws.cell(2, bl_idx).alignment.wrap_text is not True, "Principal não deve quebrar texto")
    sms = wb["sms"]
    sms_h = [c.value for c in sms[1]]
    _assert("TELEFONE" in sms_h and "Contato" in sms_h and "Nome" in sms_h, sms_h)
    _assert("Blacklist" in sms_h, sms_h)
    _assert("Data_inclusao_blacklist" in sms_h, sms_h)
    _assert(sms.cell(2, sms_h.index("TELEFONE") + 1).alignment.wrap_text is not True, "sms wrap")


def test_sms_telefone_incorreto_so_na_linha() -> None:
    bl = _bl_vazio()
    bl["TELEFONE"].add("11999990000")
    motivo = {("TELEFONE", "11999990000"): "Telefone Incorreto"}
    rec = _linha_processo(
        {
            "processo_id": 1,
            "cpf": "12345678909",
            "pessoa_nome": "Fulano",
            "numero_processo": "PROC-1",
            "numero_incidente": "1",
            "requerente": "Fulano",
            "telefones_sms": "11999990000; 11888887777",
            "telefones_hsm": "",
            "emails": "",
        },
        bl,
        motivo,
    )
    _assert(rec["Blacklist"] == "Não", rec)
    _assert("Telefone Incorreto" not in (rec.get("Motivo_blacklist") or ""), rec)
    ruim = status_linha_sms(
        cpf=rec["CPF"],
        nome="Fulano",
        requerente="Fulano",
        processo="PROC-1",
        incidente="1",
        telefone="11999990000",
        bl=bl,
        motivo_map=motivo,
    )
    bom = status_linha_sms(
        cpf=rec["CPF"],
        nome="Fulano",
        requerente="Fulano",
        processo="PROC-1",
        incidente="1",
        telefone="11888887777",
        bl=bl,
        motivo_map=motivo,
    )
    _assert(ruim["Blacklist"] == "Sim", ruim)
    _assert("Telefone Incorreto" in ruim["Motivo_blacklist"], ruim)
    _assert(bom["Blacklist"] == "Não", bom)

    buf = gerar_excel_enriquecimento(
        [rec],
        [],
        {"Processos_enriquecidos": 1},
    )
    import openpyxl

    wb = openpyxl.load_workbook(buf)
    sms = wb["sms"]
    headers = [c.value for c in sms[1]]
    tel_i = headers.index("TELEFONE") + 1
    bl_i = headers.index("Blacklist") + 1
    mot_i = headers.index("Motivo_blacklist") + 1
    by_tel = {}
    for row in sms.iter_rows(min_row=2, values_only=True):
        by_tel[str(row[tel_i - 1])] = (row[bl_i - 1], row[mot_i - 1])
    _assert(by_tel["11999990000"][0] == "Sim", by_tel)
    _assert("Telefone Incorreto" in str(by_tel["11999990000"][1]), by_tel)
    _assert(by_tel["11888887777"][0] == "Não", by_tel)


def test_sms_outra_tag_em_todas_linhas() -> None:
    bl = _bl_vazio()
    bl["TELEFONE"].add("11999990000")
    bl["CPF"].add("12345678909")
    motivo = {
        ("TELEFONE", "11999990000"): "Solicitou remoção",
        ("CPF", "12345678909"): "Sem Interesse",
    }
    rec = _linha_processo(
        {
            "processo_id": 1,
            "cpf": "12345678909",
            "pessoa_nome": "Fulano",
            "numero_processo": "PROC-1",
            "numero_incidente": "1",
            "requerente": "Fulano",
            "telefones_sms": "11999990000; 11888887777",
            "telefones_hsm": "",
            "emails": "",
        },
        bl,
        motivo,
    )
    for tel in ("11999990000", "11888887777"):
        st = rec["_sms_status"][tel]
        _assert(st["Blacklist"] == "Sim", (tel, st))
        _assert("Solicitou remoção" in st["Motivo_blacklist"], (tel, st))
        _assert("Sem Interesse" in st["Motivo_blacklist"], (tel, st))


def test_sms_misto_incorreto_e_outra_tag() -> None:
    bl = _bl_vazio()
    bl["TELEFONE"].add("11999990000")
    bl["CPF"].add("12345678909")
    motivo = {
        ("TELEFONE", "11999990000"): "Telefone Incorreto",
        ("CPF", "12345678909"): "Sem Interesse",
    }
    datas = {
        ("TELEFONE", "11999990000"): datetime(2026, 9, 1, 10, 0),
        ("CPF", "12345678909"): datetime(2026, 8, 15, 9, 0),
    }
    rec = _linha_processo(
        {
            "processo_id": 1,
            "cpf": "12345678909",
            "pessoa_nome": "Fulano",
            "numero_processo": "PROC-1",
            "numero_incidente": "1",
            "requerente": "Fulano",
            "telefones_sms": "11999990000; 11888887777",
            "telefones_hsm": "",
            "emails": "",
        },
        bl,
        motivo,
        datas,
    )
    ruim = rec["_sms_status"]["11999990000"]
    bom = rec["_sms_status"]["11888887777"]
    _assert(ruim["Blacklist"] == "Sim", ruim)
    _assert("Telefone Incorreto" in ruim["Motivo_blacklist"], ruim)
    _assert("Sem Interesse" in ruim["Motivo_blacklist"], ruim)
    _assert(bom["Blacklist"] == "Sim", bom)
    _assert("Sem Interesse" in bom["Motivo_blacklist"], bom)
    _assert("Telefone Incorreto" not in bom["Motivo_blacklist"], bom)
    _assert("15/08/2026 09:00" in bom["Data_inclusao_blacklist"], bom)
    _assert("01/09/2026 10:00" not in bom["Data_inclusao_blacklist"], bom)
    _assert("15/08/2026 09:00" in ruim["Data_inclusao_blacklist"], ruim)
    _assert("01/09/2026 10:00" in ruim["Data_inclusao_blacklist"], ruim)


if __name__ == "__main__":
    test_csv_coluna_cpf()
    test_csv_so_uma_coluna()
    test_csv_ponto_virgula()
    test_csv_vazio()
    test_blacklist_sim_motivo()
    test_blacklist_nao()
    test_excel_tem_coluna_blacklist()
    test_sms_telefone_incorreto_so_na_linha()
    test_sms_outra_tag_em_todas_linhas()
    test_sms_misto_incorreto_e_outra_tag()
    print("OK test_enriquecimento_cpf")
